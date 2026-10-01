from datetime import date
from io import BytesIO

from flask import (
    Blueprint,
    render_template,
    request,
    send_file
)

from flask_login import login_required

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle
)

from app import db
from app.models import (
    Attendance,
    ClassSession,
    Student,
    Subject
)


reports_bp = Blueprint(
    "reports",
    __name__,
    url_prefix="/reports"
)


# =========================================================
# REPORT DATA
# =========================================================

def get_report_data():

    from_date = request.args.get(
        "from_date",
        ""
    ).strip()

    to_date = request.args.get(
        "to_date",
        ""
    ).strip()

    selected_from_date = None
    selected_to_date = None

    # -----------------------------------------------------
    # PARSE FROM DATE
    # -----------------------------------------------------

    if from_date:

        try:

            selected_from_date = date.fromisoformat(
                from_date
            )

        except ValueError:

            selected_from_date = None

    # -----------------------------------------------------
    # PARSE TO DATE
    # -----------------------------------------------------

    if to_date:

        try:

            selected_to_date = date.fromisoformat(
                to_date
            )

        except ValueError:

            selected_to_date = None

    # -----------------------------------------------------
    # LOAD ALL SESSIONS IN DATE RANGE
    # -----------------------------------------------------

    session_query = db.select(
        ClassSession
    )

    if selected_from_date:

        session_query = session_query.where(
            ClassSession.session_date
            >= selected_from_date
        )

    if selected_to_date:

        session_query = session_query.where(
            ClassSession.session_date
            <= selected_to_date
        )

    all_sessions = db.session.execute(
        session_query
    ).scalars().all()

    # -----------------------------------------------------
    # SEPARATE FINALIZED AND PENDING SESSIONS
    # -----------------------------------------------------

    finalized_sessions = [
        session
        for session in all_sessions
        if session.attendance_finalized
    ]

    pending_sessions = [
        session
        for session in all_sessions
        if not session.attendance_finalized
    ]

    total_sessions = len(
        finalized_sessions
    )

    pending_session_count = len(
        pending_sessions
    )

    finalized_session_ids = [
        session.id
        for session in finalized_sessions
    ]

    # -----------------------------------------------------
    # STUDENT COUNT
    # -----------------------------------------------------

    student_count = db.session.execute(
        db.select(
            db.func.count(Student.id)
        )
    ).scalar_one()

    # -----------------------------------------------------
    # ATTENDANCE RECORDS
    #
    # IMPORTANT:
    # Only records belonging to finalized sessions are
    # included in reports.
    # -----------------------------------------------------

    attendance_records = []

    if finalized_session_ids:

        attendance_records = db.session.execute(
            db.select(
                Attendance
            ).where(
                Attendance.session_id.in_(
                    finalized_session_ids
                )
            )
        ).scalars().all()

    # -----------------------------------------------------
    # OVERALL ATTENDANCE
    # -----------------------------------------------------

    present_count = sum(
        1
        for record in attendance_records
        if record.status in (
            "present",
            "late"
        )
    )

    absent_count = sum(
        1
        for record in attendance_records
        if record.status == "absent"
    )

    total_possible_attendance = (
        student_count
        * total_sessions
    )

    if total_possible_attendance > 0:

        overall_percentage = round(
            (
                present_count
                / total_possible_attendance
            )
            * 100,
            1
        )

    else:

        overall_percentage = 0

    # -----------------------------------------------------
    # STUDENT REPORT
    # -----------------------------------------------------

    students = db.session.execute(
        db.select(Student)
        .order_by(
            Student.roll_number
        )
    ).scalars().all()

    student_reports = []

    for student in students:

        attended = sum(
            1
            for record in attendance_records
            if (
                record.student_id
                == student.id
                and record.status
                in ("present", "late")
            )
        )

        absent = sum(
            1
            for record in attendance_records
            if (
                record.student_id
                == student.id
                and record.status
                == "absent"
            )
        )

        if total_sessions > 0:

            percentage = round(
                (
                    attended
                    / total_sessions
                )
                * 100,
                1
            )

        else:

            percentage = 0

        student_reports.append(
            {
                "student": student,
                "attended": attended,
                "absent": absent,
                "total_sessions": total_sessions,
                "percentage": percentage
            }
        )

    # -----------------------------------------------------
    # SUBJECT REPORT
    # -----------------------------------------------------

    subjects = db.session.execute(
        db.select(Subject)
        .order_by(
            Subject.subject_code
        )
    ).scalars().all()

    subject_reports = []

    for subject in subjects:

        subject_sessions = [
            session
            for session in finalized_sessions
            if session.subject_id
            == subject.id
        ]

        subject_session_ids = [
            session.id
            for session in subject_sessions
        ]

        subject_attendance = [
            record
            for record in attendance_records
            if record.session_id
            in subject_session_ids
        ]

        subject_attended = sum(
            1
            for record in subject_attendance
            if record.status in (
                "present",
                "late"
            )
        )

        subject_absent = sum(
            1
            for record in subject_attendance
            if record.status == "absent"
        )

        subject_possible = (
            student_count
            * len(subject_sessions)
        )

        if subject_possible > 0:

            percentage = round(
                (
                    subject_attended
                    / subject_possible
                )
                * 100,
                1
            )

        else:

            percentage = 0

        subject_reports.append(
            {
                "subject": subject,
                "sessions": len(
                    subject_sessions
                ),
                "attended": subject_attended,
                "absent": subject_absent,
                "possible": subject_possible,
                "percentage": percentage
            }
        )

    return {
        "from_date": from_date,
        "to_date": to_date,

        "student_count": student_count,

        # Finalized sessions only
        "total_sessions": total_sessions,

        # Sessions awaiting finalization
        "pending_sessions": pending_session_count,

        # Useful for the UI if needed
        "total_scheduled_sessions": len(
            all_sessions
        ),

        "present_count": present_count,
        "absent_count": absent_count,

        "overall_percentage": overall_percentage,

        "student_reports": student_reports,
        "subject_reports": subject_reports
    }


# =========================================================
# REPORT PAGE
# =========================================================

@reports_bp.route("/")
@login_required
def reports():

    report_data = get_report_data()

    return render_template(
        "reports/index.html",
        **report_data
    )


# =========================================================
# EXCEL EXPORT
# =========================================================

@reports_bp.route("/export/excel")
@login_required
def export_excel():

    report_data = get_report_data()

    workbook = Workbook()

    # =====================================================
    # SUMMARY SHEET
    # =====================================================

    summary_sheet = workbook.active

    summary_sheet.title = "Summary"

    summary_sheet["A1"] = (
        "SmartAttend Attendance Report"
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        size=16
    )

    summary_sheet["A3"] = "Report Period"

    summary_sheet["B3"] = (
        f"{report_data['from_date'] or 'All'} "
        f"to "
        f"{report_data['to_date'] or 'All'}"
    )

    summary_sheet["A5"] = "Total Students"

    summary_sheet["B5"] = (
        report_data["student_count"]
    )

    summary_sheet["A6"] = "Finalized Class Sessions"

    summary_sheet["B6"] = (
        report_data["total_sessions"]
    )

    summary_sheet["A7"] = "Pending Sessions"

    summary_sheet["B7"] = (
        report_data["pending_sessions"]
    )

    summary_sheet["A8"] = "Attendance Marked"

    summary_sheet["B8"] = (
        report_data["present_count"]
    )

    summary_sheet["A9"] = "Absent Records"

    summary_sheet["B9"] = (
        report_data["absent_count"]
    )

    summary_sheet["A10"] = "Overall Attendance"

    summary_sheet["B10"] = (
        f"{report_data['overall_percentage']}%"
    )

    # =====================================================
    # STUDENT SHEET
    # =====================================================

    student_sheet = workbook.create_sheet(
        "Student Attendance"
    )

    student_headers = [
        "Student",
        "Roll Number",
        "Department",
        "Classes Attended",
        "Classes Absent",
        "Classes Conducted",
        "Attendance %"
    ]

    student_sheet.append(
        student_headers
    )

    for cell in student_sheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for report in report_data[
        "student_reports"
    ]:

        student = report["student"]

        student_sheet.append(
            [
                student.name,
                student.roll_number,
                student.department or "",
                report["attended"],
                report["absent"],
                report["total_sessions"],
                report["percentage"]
            ]
        )

    # =====================================================
    # SUBJECT SHEET
    # =====================================================

    subject_sheet = workbook.create_sheet(
        "Subject Attendance"
    )

    subject_headers = [
        "Subject Code",
        "Subject Name",
        "Classes Conducted",
        "Attendance Records",
        "Absent Records",
        "Attendance %"
    ]

    subject_sheet.append(
        subject_headers
    )

    for cell in subject_sheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for report in report_data[
        "subject_reports"
    ]:

        subject = report["subject"]

        subject_sheet.append(
            [
                subject.subject_code,
                subject.subject_name,
                report["sessions"],
                report["attended"],
                report["absent"],
                report["percentage"]
            ]
        )

    # =====================================================
    # AUTO COLUMN WIDTH
    # =====================================================

    for sheet in workbook.worksheets:

        for column_cells in sheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    cell_length = len(
                        str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        cell_length
                    )

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                40
            )

    # =====================================================
    # SAVE EXCEL
    # =====================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=(
            "SmartAttend_Attendance_Report.xlsx"
        ),
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )


# =========================================================
# PDF EXPORT
# =========================================================

@reports_bp.route("/export/pdf")
@login_required
def export_pdf():

    report_data = get_report_data()

    # -----------------------------------------------------
    # CREATE PDF IN MEMORY
    # -----------------------------------------------------

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm
    )

    # -----------------------------------------------------
    # STYLES
    # -----------------------------------------------------

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=20,
        leading=24,
        spaceAfter=8
    )

    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=10,
        leading=14,
        textColor=colors.grey,
        spaceAfter=20
    )

    section_style = ParagraphStyle(
        "SectionTitle",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        spaceBefore=12,
        spaceAfter=8
    )

    story = []

    # =====================================================
    # TITLE
    # =====================================================

    story.append(
        Paragraph(
            "SmartAttend Attendance Report",
            title_style
        )
    )

    report_period = (
        f"{report_data['from_date'] or 'All dates'} "
        f"to "
        f"{report_data['to_date'] or 'All dates'}"
    )

    story.append(
        Paragraph(
            f"Report Period: {report_period}",
            subtitle_style
        )
    )

    # =====================================================
    # SUMMARY
    # =====================================================

    story.append(
        Paragraph(
            "Attendance Summary",
            section_style
        )
    )

    summary_data = [
        [
            "Metric",
            "Value"
        ],
        [
            "Total Students",
            str(
                report_data["student_count"]
            )
        ],
        [
            "Finalized Class Sessions",
            str(
                report_data["total_sessions"]
            )
        ],
        [
            "Pending Sessions",
            str(
                report_data["pending_sessions"]
            )
        ],
        [
            "Attendance Marked",
            str(
                report_data["present_count"]
            )
        ],
        [
            "Absent Records",
            str(
                report_data["absent_count"]
            )
        ],
        [
            "Overall Attendance",
            f"{report_data['overall_percentage']}%"
        ]
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            90 * mm,
            70 * mm
        ]
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#212529")
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "FONTNAME",
                    (0, 1),
                    (-1, -1),
                    "Helvetica"
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),

                (
                    "ALIGN",
                    (1, 1),
                    (1, -1),
                    "CENTER"
                ),

                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    7
                ),

                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    7
                )
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(
            1,
            12
        )
    )

    # =====================================================
    # STUDENT ATTENDANCE
    # =====================================================

    story.append(
        Paragraph(
            "Student-wise Attendance",
            section_style
        )
    )

    student_data = [
        [
            "Student",
            "Roll No.",
            "Department",
            "Attended",
            "Absent",
            "Conducted",
            "Attendance"
        ]
    ]

    for report in report_data[
        "student_reports"
    ]:

        student = report["student"]

        student_data.append(
            [
                student.name,
                student.roll_number,
                student.department or "-",
                str(
                    report["attended"]
                ),
                str(
                    report["absent"]
                ),
                str(
                    report["total_sessions"]
                ),
                f"{report['percentage']}%"
            ]
        )

    student_table = Table(
        student_data,
        colWidths=[
            34 * mm,
            22 * mm,
            28 * mm,
            20 * mm,
            20 * mm,
            22 * mm,
            24 * mm
        ],
        repeatRows=1
    )

    student_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#212529")
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.5
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),

                (
                    "ALIGN",
                    (3, 1),
                    (-1, -1),
                    "CENTER"
                ),

                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                ),

                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                )
            ]
        )
    )

    story.append(
        student_table
    )

    story.append(
        Spacer(
            1,
            15
        )
    )

    # =====================================================
    # SUBJECT ATTENDANCE
    # =====================================================

    story.append(
        Paragraph(
            "Subject-wise Attendance",
            section_style
        )
    )

    subject_data = [
        [
            "Subject Code",
            "Subject Name",
            "Sessions",
            "Attendance",
            "Absent",
            "Percentage"
        ]
    ]

    for report in report_data[
        "subject_reports"
    ]:

        subject = report["subject"]

        subject_data.append(
            [
                subject.subject_code,
                subject.subject_name,
                str(
                    report["sessions"]
                ),
                str(
                    report["attended"]
                ),
                str(
                    report["absent"]
                ),
                f"{report['percentage']}%"
            ]
        )

    subject_table = Table(
        subject_data,
        colWidths=[
            27 * mm,
            60 * mm,
            22 * mm,
            27 * mm,
            22 * mm,
            27 * mm
        ],
        repeatRows=1
    )

    subject_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#212529")
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),

                (
                    "ALIGN",
                    (2, 1),
                    (-1, -1),
                    "CENTER"
                ),

                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                ),

                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5
                )
            ]
        )
    )

    story.append(
        subject_table
    )

    # =====================================================
    # BUILD PDF
    # =====================================================

    document.build(
        story
    )

    output.seek(0)

    # =====================================================
    # SEND PDF
    # =====================================================

    return send_file(
        output,
        as_attachment=True,
        download_name=(
            "SmartAttend_Attendance_Report.pdf"
        ),
        mimetype="application/pdf"
    )